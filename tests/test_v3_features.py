"""Tests for v3 Career OS features:

  recommendation engine, research caching, document versioning, analytics,
  pipeline run persistence, multi-sheet dashboard, event logging, and the
  recommendation dashboard.
"""
import sqlite3
import tempfile
import unittest
from contextlib import contextmanager
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

import openpyxl

from agents import document_agent, research_agent
from core import analytics, database, event_log, exporter, pipeline, pipeline_run
from core.recommendation_engine import (
    RECOMMENDATION_ORDER,
    recommend,
    recommend_all,
    sort_for_dashboard,
)


# ── 1. Recommendation engine ─────────────────────────────────────────────────────

class TestRecommendationEngine(unittest.TestCase):
    def _job(self, **kw):
        base = {"id": 1, "total_score": 95, "role_category": "ai_strategy",
                "visa_gate": "PASS", "location": "Milan, Italy", "work_mode": "Hybrid",
                "application_status": "Not Started"}
        base.update(kw)
        return base

    def test_apply_immediately(self):
        out = recommend(self._job(total_score=95), has_research=True)
        self.assertEqual(out["recommendation"], "Apply Immediately")
        self.assertTrue(out["recommendation_reason"])

    def test_apply_this_week(self):
        out = recommend(self._job(total_score=75), has_research=True)
        self.assertEqual(out["recommendation"], "Apply This Week")

    def test_research_company_when_unresearched(self):
        out = recommend(self._job(total_score=88), has_research=False)
        self.assertEqual(out["recommendation"], "Research Company")

    def test_find_referral_when_no_sponsorship(self):
        out = recommend(self._job(total_score=92, visa_gate="FAIL"), has_research=True)
        self.assertEqual(out["recommendation"], "Find Referral First")

    def test_save_for_later(self):
        out = recommend(self._job(total_score=60))
        self.assertEqual(out["recommendation"], "Save For Later")

    def test_skip_low_score(self):
        out = recommend(self._job(total_score=30))
        self.assertEqual(out["recommendation"], "Skip")

    def test_skip_when_already_in_pipeline(self):
        out = recommend(self._job(total_score=95, application_status="Interview"))
        self.assertEqual(out["recommendation"], "Skip")
        self.assertIn("Interview", out["recommendation_reason"][0])

    def test_reasons_include_signals(self):
        out = recommend(self._job(total_score=95), has_research=True)
        self.assertIn("High sponsorship likelihood", out["recommendation_reason"])
        self.assertIn("AI initiatives align with MBA specialization",
                      out["recommendation_reason"])

    def test_recommend_all_does_not_mutate(self):
        rows = [self._job(id=1, total_score=95)]
        recommend_all(rows, research_job_ids={1})
        self.assertNotIn("recommendation", rows[0])

    def test_sort_for_dashboard_orders_by_recommendation_then_score(self):
        rows = recommend_all([
            self._job(id=1, total_score=72, visa_gate="PASS"),     # Apply This Week
            self._job(id=2, total_score=95, visa_gate="PASS"),     # Apply Immediately
        ], research_job_ids={1, 2})
        ordered = sort_for_dashboard(rows)
        self.assertEqual(ordered[0]["id"], 2)
        self.assertLessEqual(RECOMMENDATION_ORDER[ordered[0]["recommendation"]],
                             RECOMMENDATION_ORDER[ordered[1]["recommendation"]])


# ── 2. Research caching ──────────────────────────────────────────────────────────

class TestResearchCaching(unittest.TestCase):
    def test_stale_when_missing(self):
        with patch.object(research_agent.database, "get_research_for_job", return_value=None):
            self.assertTrue(research_agent.research_is_stale(1))

    def test_fresh_within_ttl(self):
        recent = (datetime.now() - timedelta(days=5)).isoformat()
        with patch.object(research_agent.database, "get_research_for_job",
                          return_value={"researched_at": recent}):
            self.assertFalse(research_agent.research_is_stale(1))

    def test_stale_beyond_ttl(self):
        old = (datetime.now() - timedelta(days=40)).isoformat()
        with patch.object(research_agent.database, "get_research_for_job",
                          return_value={"researched_at": old}):
            self.assertTrue(research_agent.research_is_stale(1))


# ── 3. Document versioning ───────────────────────────────────────────────────────

class TestDocumentVersioning(unittest.TestCase):
    def _job(self):
        return {"id": 1, "title": "AI PM", "company": "Acme", "description": "Build AI."}

    def test_reuses_when_hashes_match(self):
        existing = {"version": 1, "source_hash": "S", "prompt_hash": "P", "research_hash": "R"}
        with patch.object(document_agent.database, "get_latest_document", return_value=existing), \
             patch.object(document_agent, "version", return_value="P"), \
             patch.object(document_agent, "_call_claude") as call:
            result = document_agent._generate_one(
                self._job(), "resume", "resume_prompt", {}, "ctx",
                source_hash="S", research_hash="R")
        self.assertEqual(result, "reused")
        call.assert_not_called()

    def test_generates_new_version_when_changed(self):
        existing = {"version": 2, "source_hash": "OLD", "prompt_hash": "P", "research_hash": "R"}
        captured = {}
        with tempfile.TemporaryDirectory() as td, \
             patch.object(document_agent.config, "OUTPUTS_DIR", td), \
             patch.object(document_agent.database, "get_latest_document", return_value=existing), \
             patch.object(document_agent, "version", return_value="P"), \
             patch.object(document_agent, "_call_claude", return_value="# New resume"), \
             patch.object(document_agent.database, "insert_document",
                          side_effect=lambda **kw: captured.update(kw)):
            result = document_agent._generate_one(
                self._job(), "resume", "resume_prompt", {}, "ctx",
                source_hash="NEW", research_hash="R")
        self.assertEqual(result, "generated")
        self.assertEqual(captured["version"], 3)             # bumped from 2
        self.assertTrue(captured["file_name"].endswith("_v3.md"))


# ── 4. Analytics engine ──────────────────────────────────────────────────────────

class TestAnalytics(unittest.TestCase):
    def setUp(self):
        self._agg = patch.object(database, "get_aggregate_counts", return_value={
            "jobs_found": 200, "jobs_scored": 150, "researched": 20, "documents": 10})
        self._status = patch.object(database, "get_status_counts", return_value={
            "Applied": 5, "Interview": 3, "Final Round": 1, "Offer": 1, "Rejected": 2})
        self._agg.start()
        self._status.start()

    def tearDown(self):
        self._agg.stop()
        self._status.stop()

    def test_application_metrics_cumulative_funnel(self):
        m = analytics.application_metrics()
        self.assertEqual(m["jobs_found"], 200)
        # Applied counts everyone at Applied or later forward stage: 5+3+1+1 = 10
        self.assertEqual(m["applied"], 10)
        self.assertEqual(m["interviews"], 5)   # 3+1+1
        self.assertEqual(m["offers"], 1)
        self.assertEqual(m["rejections"], 2)

    def test_conversion_metrics(self):
        c = analytics.conversion_metrics()
        self.assertEqual(c["conversion_rate"], round(100 * 10 / 150, 1))
        self.assertEqual(c["interview_rate"], round(100 * 5 / 10, 1))
        self.assertEqual(c["offer_rate"], round(100 * 1 / 10, 1))

    def test_funnel_and_status_breakdown(self):
        funnel = analytics.funnel_metrics()
        self.assertEqual(funnel[0], {"stage": "Jobs Found", "count": 200})
        sb = analytics.status_breakdown()
        self.assertEqual(sb["Applied"], 5)
        self.assertIn("Not Started", sb)  # zero-filled lifecycle

    def test_zero_division_safe(self):
        with patch.object(database, "get_status_counts", return_value={}):
            c = analytics.conversion_metrics()
            self.assertEqual(c["interview_rate"], 0.0)


# ── 5. Pipeline run persistence ──────────────────────────────────────────────────

def _runs_db():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.execute("""
        CREATE TABLE pipeline_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT, run_id TEXT UNIQUE,
            started_at TEXT, completed_at TEXT, jobs_found INTEGER, jobs_scored INTEGER,
            companies_researched INTEGER, documents_generated INTEGER,
            applications_updated INTEGER, status TEXT, duration_seconds REAL, detail TEXT
        )
    """)
    return conn


class TestPipelineRunPersistence(unittest.TestCase):
    def setUp(self):
        self.conn = _runs_db()

        @contextmanager
        def _ctx():
            yield self.conn
            self.conn.commit()

        self._patch = patch("core.database.get_connection", _ctx)
        self._patch.start()

    def tearDown(self):
        self._patch.stop()
        self.conn.close()

    def test_start_finalize_save_and_read(self):
        run = pipeline_run.PipelineRun.start()
        run.jobs_found = 200
        run.finalize(pipeline_run.PipelineRun.SUCCESS)
        self.assertIsNotNone(run.duration_seconds)
        run.save()
        rows = pipeline_run.recent()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["status"], "SUCCESS")
        self.assertEqual(rows[0]["jobs_found"], 200)

    def test_save_is_idempotent_on_run_id(self):
        run = pipeline_run.PipelineRun.start()
        run.finalize(pipeline_run.PipelineRun.PARTIAL_SUCCESS)
        run.save()
        run.jobs_scored = 99
        run.save()
        rows = pipeline_run.recent()
        self.assertEqual(len(rows), 1)          # upsert, not duplicate
        self.assertEqual(rows[0]["jobs_scored"], 99)


# ── 6. Multi-sheet dashboard ─────────────────────────────────────────────────────

_SHEET_ROWS = [
    {"id": 1, "title": "AI PM", "company": "Acme", "location": "Milan, Italy",
     "description": "remote ai", "role_category": "ai_strategy", "url": "http://a/1",
     "fetched_date": "2026-06-10", "visa_gate": "PASS", "total_score": 95,
     "application_status": "Not Started"},
    {"id": 2, "title": "Strategy", "company": "Beta", "location": "London",
     "description": "hybrid", "role_category": "business_strategy", "url": "http://b/2",
     "fetched_date": "2026-06-11", "visa_gate": "PASS", "total_score": 82,
     "application_status": "Applied"},
]

_ANALYTICS_SUMMARY = {
    "application_metrics": {"jobs_found": 2, "jobs_scored": 2, "researched": 0,
                           "documents": 0, "applied": 1, "assessments": 0,
                           "interviews": 0, "final_rounds": 0, "offers": 0, "rejections": 0},
    "conversion_metrics": {"conversion_rate": 50.0, "interview_rate": 0.0, "offer_rate": 0.0},
    "funnel_metrics": [{"stage": "Jobs Found", "count": 2}, {"stage": "Jobs Scored", "count": 2},
                       {"stage": "Applied", "count": 1}, {"stage": "Assessments", "count": 0},
                       {"stage": "Interviews", "count": 0}, {"stage": "Final Rounds", "count": 0},
                       {"stage": "Offers", "count": 0}],
    "status_breakdown": {"Not Started": 1, "Applied": 1},
}


class TestMultiSheetDashboard(unittest.TestCase):
    def test_all_sheets_present_with_chart(self):
        with tempfile.TemporaryDirectory() as td:
            out = str(Path(td) / "jobs_master.xlsx")
            with patch.object(database, "get_all_scored_jobs_ranked", return_value=_SHEET_ROWS), \
                 patch.object(database, "get_researched_job_ids", return_value=set()), \
                 patch("core.tracker.list_applications", return_value=[]), \
                 patch("core.analytics.summary", return_value=_ANALYTICS_SUMMARY):
                exporter.export_jobs_master_xlsx(out)
            wb = openpyxl.load_workbook(out)
            for sheet in ("Jobs Master", "Apply Now", "This Week", "Applications", "Analytics"):
                self.assertIn(sheet, wb.sheetnames)
            # Apply Now has the >=90 job (Acme), This Week has the 80-89 job (Beta).
            self.assertEqual(wb["Apply Now"].cell(row=2, column=4).value, "Acme")
            self.assertEqual(wb["This Week"].cell(row=2, column=4).value, "Beta")
            self.assertGreaterEqual(len(wb["Analytics"]._charts), 1)


# ── 7. Event logging ─────────────────────────────────────────────────────────────

class TestEventLogging(unittest.TestCase):
    @staticmethod
    def _reset_loggers():
        """Detach + close any file handlers so a temp-dir path never leaks into
        later tests (logging keeps loggers globally, not just our cache)."""
        for lg in event_log._loggers.values():
            for h in list(lg.handlers):
                h.close()
                lg.removeHandler(h)
        event_log._loggers.clear()

    def test_writes_structured_line(self):
        self._reset_loggers()
        try:
            with tempfile.TemporaryDirectory() as td:
                with patch.object(event_log.config, "LOGS_DIR", td):
                    event_log.log_event("pipeline", "step_done", job_id=7,
                                         company="Acme", duration=1.25, status="ok")
                    log_file = Path(td) / "pipeline.log"
                    self.assertTrue(log_file.exists())
                    content = log_file.read_text()
        finally:
            self._reset_loggers()
        self.assertIn("event=step_done", content)
        self.assertIn("job_id=7", content)
        self.assertIn("status=ok", content)


# ── 8. Recommendation dashboard ──────────────────────────────────────────────────

class TestRecommendationDashboard(unittest.TestCase):
    def test_filters_and_sorts(self):
        rows = [dict(r) for r in _SHEET_ROWS] + [
            {"id": 3, "title": "Low", "company": "Gamma", "location": "Berlin",
             "role_category": "hr_analytics", "url": "x", "fetched_date": "2026-06-01",
             "visa_gate": None, "total_score": 40, "application_status": "Not Started"},
        ]
        with patch.object(database, "get_all_scored_jobs_ranked", return_value=rows), \
             patch.object(database, "get_researched_job_ids", return_value={1, 2}), \
             patch("core.profile_loader.load", return_value={}):
            shown = pipeline.display_recommendations(limit=20, min_score=70)
        # Only the two >=70 jobs survive the default filter.
        self.assertEqual(len(shown), 2)
        self.assertTrue(all("recommendation" in r for r in shown))


if __name__ == "__main__":
    unittest.main(verbosity=2)
