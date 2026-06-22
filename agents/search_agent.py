import hashlib
import html
import json
import os
import re
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

import requests
from langdetect import DetectorFactory, detect_langs
from langdetect.lang_detect_exception import LangDetectException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import config, database
from core.profile_loader import load as load_profile
from core.role_loader import classify_title

DetectorFactory.seed = 0

# ── Track membership ───────────────────────────────────────────────────────────

_TRACK_A = {"product_management", "ai_strategy", "business_strategy", "digital_transformation"}
_TRACK_B = {"people_analytics", "hr_analytics", "workforce_planning", "talent_acquisition"}

# ── Seniority patterns ─────────────────────────────────────────────────────────

_HARD_REJECT_TITLE = re.compile(
    r"\b(director|head of|chief|vice president|distinguished|c-level|cxo|president)\b",
    re.IGNORECASE,
)
_VP_REJECT = re.compile(r"\bvp\b", re.IGNORECASE)
_SOFT_REJECT_TITLE = re.compile(r"\b(senior|sr\.?|lead)\b", re.IGNORECASE)
_PARTNER_REJECT = re.compile(r"\bpartner\b", re.IGNORECASE)
_ACCEPT_SIGNALS = re.compile(
    r"\b(intern|internship|stage|stagista|stagiaire|praktikum|praktikant"
    r"|tirocinio|tirocinante|junior|jr\.?|graduate|associate|fellowship"
    r"|trainee|rotational|entry.?level|analyst|apprenti)\b",
    re.IGNORECASE,
)
_ALWAYS_ACCEPT_TITLE = re.compile(
    r"\b(product manager|hr analytics manager|people analytics manager"
    r"|workforce planning manager|business analyst|hr data analyst"
    r"|people data analyst|hr reporting manager|talent analytics manager)\b",
    re.IGNORECASE,
)
_DESC_OVERQUALIFIED = re.compile(
    r"\b([5-9]|1[0-9])\+\s*years?\s*(?:of\s+)?(?:experience|exp)\b",
    re.IGNORECASE,
)

# ── Seniority filter — unified for Track A and Track B ────────────────────────
# Only internship-style roles accepted. Accept always overrides reject in same title.

# Accept signals: intern / graduate / trainee / working-student / stage / praktikum / …
_INTERN_ACCEPT = re.compile(
    r"\b(intern|internship|working[\s\-]?student|werkstudent|graduate(?:\s+program)?"
    r"|trainee|apprentice|apprenticeship|stage|stagiaire|praktikum|praktikant"
    r"|tirocinio|tirocinante)\b",
    re.IGNORECASE,
)

# Reject signals: role-type/seniority indicators that disqualify a title
# Only fires when no accept signal is present.
_SENIORITY_REJECT = re.compile(
    r"\b(analyst|associate|specialist|coordinator|junior|manager|lead|principal"
    r"|director|head(?:\s+of)?|vp|vice[\s\-]?president|chief|partner)\b",
    re.IGNORECASE,
)

# ── Geography patterns ─────────────────────────────────────────────────────────

_NON_EU = re.compile(
    r"\b(united states|usa|u\.s\.a|us|new york|nyc"
    r"|san francisco|sf|los angeles|chicago|chi|boston"
    r"|seattle|sea|austin|atlanta|atl"
    r"|canada|toronto|vancouver|montreal"
    r"|india|bangalore|bengaluru|mumbai|delhi|hyderabad|pune"
    r"|china|beijing|shanghai|singapore|tokyo|japan|hong kong"
    r"|australia|sydney|brazil|s[aã]o paulo|mexico"
    r"|dubai|uae|middle east|apac|latam|mena)\b",
    re.IGNORECASE,
)
_EU_OK = re.compile(
    r"\b(remote|hybrid|europe|european|anywhere|home office"
    # EU-27 + UK + Switzerland — country names
    r"|italy|italia|uk|united kingdom|germany|deutschland|france"
    r"|netherlands|spain|ireland|switzerland|belgium|sweden|denmark"
    r"|finland|norway|poland|austria|portugal|czech(?:ia)?|hungary|romania"
    r"|bulgaria|lithuania|latvia|estonia|greece|cyprus|croatia"
    r"|slovakia|slovenia|luxembourg|malta"
    # Major cities — original set
    r"|milan|milano|london|amsterdam|berlin|paris|barcelona|dublin|zurich"
    r"|rome|roma|munich|münchen|lyon|marseille|madrid|vienna|wien"
    r"|stockholm|copenhagen|oslo|helsinki|brussels|warsaw|lisbon|zagreb"
    # Cities identified from unknown_reject audit
    r"|m[aá]laga|d[uü]sseldorf|hamburg|frankfurt|budapest|sofia|vilnius"
    r"|riga|athens|antwerp|nantes|le[oó]n|nicosia|edinburgh|crawley"
    # Additional major EU/UK/CH hiring hubs
    r"|cologne|k[oö]ln|stuttgart|tallinn|bratislava|ljubljana"
    r"|valencia|seville|sevilla|porto"
    r"|gothenburg|aarhus|rotterdam|eindhoven|the hague|den haag|utrecht"
    r"|krak[oó]w|wroc[lł]aw|gda[nń]sk"
    r"|toulouse|bordeaux|nice|glasgow|manchester|cambridge|oxford"
    r"|geneva|basel|lausanne"
    r"|bologna|turin|torino|florence|firenze)\b",
    re.IGNORECASE,
)

# ── Company normalisation for dedup hash ───────────────────────────────────────

_COMPANY_SUFFIXES = re.compile(
    r"\b(gmbh|spa|s\.p\.a\.|ltd|limited|inc|incorporated|se|ag|nv|bv"
    r"|srl|sarl|sa|plc|llc|llp|kg|co|group|holding|holdings)\b",
    re.IGNORECASE,
)

_COUNTRY_MAP = {
    "italy": "it", "italia": "it", "milan": "it", "milano": "it",
    "rome": "it", "roma": "it", "turin": "it",
    "uk": "gb", "united kingdom": "gb", "england": "gb", "london": "gb",
    "germany": "de", "deutschland": "de", "berlin": "de", "munich": "de",
    "france": "fr", "paris": "fr", "lyon": "fr",
    "netherlands": "nl", "amsterdam": "nl",
    "spain": "es", "barcelona": "es", "madrid": "es",
    "ireland": "ie", "dublin": "ie",
    "switzerland": "ch", "zurich": "ch",
    "sweden": "se", "stockholm": "se",
    "belgium": "be", "brussels": "be",
    "austria": "at", "vienna": "at",
    "remote": "eu", "europe": "eu", "european": "eu",
}


# ── Helpers ────────────────────────────────────────────────────────────────────

def _strip_html(text: str) -> str:
    if not text:
        return ""
    text = re.sub(r"<[^>]+>", " ", text)
    text = html.unescape(text)
    return re.sub(r"\s+", " ", text).strip()


def _normalize_token(s: str) -> str:
    s = s.lower()
    s = _COMPANY_SUFFIXES.sub("", s)
    s = re.sub(r"[^\w\s]", "", s)
    return re.sub(r"\s+", " ", s).strip()


def _country_code(location: str) -> str:
    if not location:
        return "eu"
    loc = location.lower()
    for kw, code in _COUNTRY_MAP.items():
        if kw in loc:
            return code
    return "eu"


def _compute_dedup_hash(company: str, title: str, location: str) -> str:
    payload = (
        _normalize_token(company) + "|"
        + _normalize_token(title) + "|"
        + _country_code(location)
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:16]


def _parse_date(raw: str, fetched: date) -> Optional[str]:
    if not raw:
        return None
    raw = str(raw).strip().lower()
    m = re.search(r"(\d+)\s*(hour|minute|day|week|month)", raw)
    if m:
        n, unit = int(m.group(1)), m.group(2)
        if "hour" in unit or "minute" in unit:
            return str(fetched)
        if "day" in unit:
            return str(fetched - timedelta(days=n))
        if "week" in unit:
            return str(fetched - timedelta(weeks=n))
        if "month" in unit:
            return str(fetched - timedelta(days=n * 30))
    try:
        return str(date.fromisoformat(raw[:10]))
    except (ValueError, TypeError):
        return None


def _detect_language(text: str) -> tuple[str, float]:
    if not text or len(text.strip()) < 50:
        return ("en", 0.0)
    try:
        results = detect_langs(text[:500])
        if results:
            top = results[0]
            return (top.lang, float(top.prob))
    except LangDetectException:
        pass
    return ("en", 0.0)


# ── Filter pipeline ────────────────────────────────────────────────────────────

def _check_exclusion(title: str, keywords: list[str]) -> tuple[bool, str]:
    tl = title.lower()
    for kw in keywords:
        if re.search(r"\b" + re.escape(kw.lower()) + r"\b", tl):
            return True, kw
    return False, ""


def _check_seniority(title: str, description: str) -> tuple[bool, str]:
    tl = title.lower()

    if _ALWAYS_ACCEPT_TITLE.search(tl):
        return True, "always_accept"

    if _HARD_REJECT_TITLE.search(tl):
        return False, "hard_reject_title"
    if _VP_REJECT.search(tl):
        return False, "vp_reject"
    if _PARTNER_REJECT.search(tl) and not _ACCEPT_SIGNALS.search(tl):
        return False, "partner_reject"

    if description and _DESC_OVERQUALIFIED.search(description[:1000]):
        return False, "overqualified_desc"

    if _SOFT_REJECT_TITLE.search(tl) and not _ACCEPT_SIGNALS.search(tl):
        return False, "soft_reject_no_accept"

    if _ACCEPT_SIGNALS.search(tl):
        return True, "accept_title"
    if description and _ACCEPT_SIGNALS.search(description[:500].lower()):
        return True, "accept_desc"

    return False, "no_accept_signal"


def _check_seniority_by_track(title: str, description: str, track: str) -> tuple[bool, str]:
    """
    Track A and Track B: accept only internship-style roles.
    Accept signal (intern/graduate/trainee/…) always overrides a reject signal present
    in the same title — e.g. 'HR Analytics Manager Intern' → accepted.
    Unclassified: delegates to generic _check_seniority() for broad compatibility.
    """
    tl = title.lower()

    if track in ("A", "B"):
        if _INTERN_ACCEPT.search(tl):
            return True, "intern_accept"
        if _SENIORITY_REJECT.search(tl):
            return False, "seniority_reject"
        return False, "no_accept_signal"

    return _check_seniority(title, description)


def _check_geography(location: str) -> tuple[bool, str]:
    if not location:
        return True, "no_location"
    if _NON_EU.search(location):
        if re.search(r"\b(anywhere|worldwide)\b", location, re.IGNORECASE):
            return True, "remote_worldwide"
        return False, "non_eu"
    if _EU_OK.search(location):
        return True, "eu"
    print(f"    [Geo] Unknown location rejected: {location!r}")
    return False, "unknown_reject"


def _check_language(lang: str, conf: float, role_cat: str) -> tuple[bool, str]:
    if conf < 0.85:
        return True, "low_confidence"
    if lang == "en":
        return True, "english"
    if lang == "it" and role_cat in _TRACK_B:
        return True, "italian_track_b"
    if lang != "en" and role_cat in _TRACK_A:
        return False, f"non_en_track_a:{lang}"
    return True, "default_accept"


def _assign_track(role_category: str) -> str:
    if role_category in _TRACK_A:
        return "A"
    if role_category in _TRACK_B:
        return "B"
    return "unclassified"


# ── Normalisers ────────────────────────────────────────────────────────────────

def _norm_greenhouse(raw: dict, fetched: date) -> dict:
    loc = raw.get("location")
    if isinstance(loc, dict):
        location = str(loc.get("name") or "")
    else:
        location = str(loc or "")
    updated = (raw.get("updated_at") or raw.get("created_at") or "")[:10] or str(fetched)
    return {
        "title":       (raw.get("title") or "").strip(),
        "company":     (raw.get("_company_name") or "").strip(),
        "location":    str(location or "").strip(),
        "job_board":   "greenhouse",
        "url":         (raw.get("absolute_url") or "").strip(),
        "description": _strip_html(raw.get("content") or ""),
        "posted_date": updated,
        "raw_data":    json.dumps(raw, ensure_ascii=False),
    }


def _norm_lever(raw: dict, fetched: date) -> dict:
    cats = raw.get("categories", {})
    if isinstance(cats, dict):
        location = cats.get("location") or cats.get("allLocations") or ""
    else:
        location = ""
    if isinstance(location, list):
        location = ", ".join(location)

    desc = raw.get("descriptionPlain") or _strip_html(raw.get("description") or "")

    created = raw.get("createdAt")
    if created:
        try:
            posted = str(datetime.fromtimestamp(int(created) / 1000).date())
        except (ValueError, TypeError):
            posted = str(fetched)
    else:
        posted = str(fetched)

    return {
        "title":       (raw.get("text") or "").strip(),
        "company":     (raw.get("_company_name") or "").strip(),
        "location":    str(location).strip(),
        "job_board":   "lever",
        "url":         (raw.get("hostedUrl") or "").strip(),
        "description": desc,
        "posted_date": posted,
        "raw_data":    json.dumps(raw, ensure_ascii=False),
    }


# ── ATS connectors ─────────────────────────────────────────────────────────────

def fetch_greenhouse(company: str, ats_id: str) -> list[dict]:
    url = f"https://boards-api.greenhouse.io/v1/boards/{ats_id}/jobs?content=true"
    try:
        resp = requests.get(url, timeout=20)
        if resp.status_code == 404:
            return []
        resp.raise_for_status()
        jobs = resp.json().get("jobs", [])
        for j in jobs:
            j["_company_name"] = company
        fetched = date.today()
        return [_norm_greenhouse(j, fetched) for j in jobs]
    except requests.RequestException as exc:
        print(f"    [Greenhouse] {company}: FAILED — {exc}")
        return []


def fetch_lever(company: str, ats_id: str) -> list[dict]:
    url = f"https://api.lever.co/v0/postings/{ats_id}?mode=json"
    try:
        resp = requests.get(url, timeout=20)
        if resp.status_code == 404:
            return []
        resp.raise_for_status()
        jobs = resp.json()
        if not isinstance(jobs, list):
            jobs = []
        for j in jobs:
            j["_company_name"] = company
        fetched = date.today()
        return [_norm_lever(j, fetched) for j in jobs]
    except requests.RequestException as exc:
        print(f"    [Lever] {company}: FAILED — {exc}")
        return []


# ── More ATS connectors (free official public APIs, no Apify needed) ───────────

def _norm_ashby(raw: dict, company: str, fetched: date) -> dict:
    loc = raw.get("location")
    if isinstance(loc, dict):
        loc = loc.get("name") or ""
    return {
        "title":       (raw.get("title") or "").strip(),
        "company":     company,
        "location":    str(loc or "").strip(),
        "job_board":   "ashby",
        "url":         (raw.get("jobUrl") or raw.get("applyUrl") or "").strip(),
        "description": _strip_html(raw.get("descriptionHtml") or raw.get("description") or ""),
        "posted_date": str(raw.get("publishedDate") or raw.get("publishedAt") or fetched)[:10],
        "raw_data":    json.dumps(raw, ensure_ascii=False),
    }


def fetch_ashby(company: str, ats_id: str) -> list[dict]:
    url = f"https://api.ashbyhq.com/posting-api/job-board/{ats_id}?includeCompensation=false"
    try:
        resp = requests.get(url, timeout=20)
        if resp.status_code == 404:
            return []
        resp.raise_for_status()
        jobs = resp.json().get("jobs", [])
        fetched = date.today()
        return [_norm_ashby(j, company, fetched) for j in jobs]
    except requests.RequestException as exc:
        print(f"    [Ashby] {company}: FAILED — {exc}")
        return []


def _norm_smartrecruiters(raw: dict, company: str, fetched: date) -> dict:
    loc = raw.get("location") or {}
    location = ", ".join(p for p in (loc.get("city"), loc.get("country")) if p)
    return {
        "title":       (raw.get("name") or "").strip(),
        "company":     company,
        "location":    location,
        "job_board":   "smartrecruiters",
        "url":         (raw.get("ref") or raw.get("applyUrl") or "").strip(),
        "description": "",  # full text needs a per-posting detail call; title gates suffice
        "posted_date": str(raw.get("releasedDate") or fetched)[:10],
        "raw_data":    json.dumps(raw, ensure_ascii=False),
    }


def fetch_smartrecruiters(company: str, ats_id: str) -> list[dict]:
    url = f"https://api.smartrecruiters.com/v1/companies/{ats_id}/postings?limit=100"
    try:
        resp = requests.get(url, timeout=20)
        if resp.status_code == 404:
            return []
        resp.raise_for_status()
        postings = resp.json().get("content", [])
        fetched = date.today()
        return [_norm_smartrecruiters(p, company, fetched) for p in postings]
    except requests.RequestException as exc:
        print(f"    [SmartRecruiters] {company}: FAILED — {exc}")
        return []


# Dispatch a watchlist company to the right ATS connector by its ats_provider.
_ATS_FETCHERS = {
    "greenhouse":      fetch_greenhouse,
    "lever":           fetch_lever,
    "ashby":           fetch_ashby,
    "smartrecruiters": fetch_smartrecruiters,
}


def fetch_company(company: dict) -> list[dict]:
    """Fetch raw jobs for one watchlist company via its ATS. [] if provider unknown."""
    fn = _ATS_FETCHERS.get(company.get("ats_provider", ""))
    return fn(company.get("name", ""), company.get("ats_id", "")) if fn else []


# ── Job processor ──────────────────────────────────────────────────────────────

def _process_job(job: dict, exclude_keywords: list[str], stats: dict) -> Optional[int]:
    title = job.get("title", "").strip()
    company = job.get("company", "").strip()
    location = job.get("location", "").strip()
    url = job.get("url", "").strip()
    description = job.get("description", "")
    source = job.get("job_board", "unknown")

    stats["fetched"] = stats.get("fetched", 0) + 1

    if not title or not company:
        stats["rejected_other"] = stats.get("rejected_other", 0) + 1
        return None

    excluded, _ = _check_exclusion(title, exclude_keywords)
    if excluded:
        stats["rejected_excluded"] = stats.get("rejected_excluded", 0) + 1
        return None

    # Classify first — track is required before the seniority gate
    role_category = classify_title(title)
    track = _assign_track(role_category)

    ok, seniority_reason = _check_seniority_by_track(title, description, track)
    if not ok:
        stats["rejected_seniority"] = stats.get("rejected_seniority", 0) + 1
        stats.setdefault("by_seniority_reason", {})[seniority_reason] = (
            stats.setdefault("by_seniority_reason", {}).get(seniority_reason, 0) + 1
        )
        return None

    ok, _ = _check_geography(location)
    if not ok:
        stats["rejected_geography"] = stats.get("rejected_geography", 0) + 1
        return None

    lang_code, lang_conf = _detect_language(description or title)

    ok, _ = _check_language(lang_code, lang_conf, role_category)
    if not ok:
        stats["rejected_language"] = stats.get("rejected_language", 0) + 1
        return None

    dedup_hash = _compute_dedup_hash(company, title, location)

    existing = database.get_job_by_hash(dedup_hash)
    if existing:
        stats["duplicates"] = stats.get("duplicates", 0) + 1
        return existing["id"]

    if url:
        existing = database.get_job_by_url(url)
        if existing:
            stats["duplicates"] = stats.get("duplicates", 0) + 1
            return existing["id"]

    try:
        raw_obj = json.loads(job.get("raw_data") or "{}")
    except (json.JSONDecodeError, TypeError):
        raw_obj = {}
    raw_obj.update({
        "_detected_language": lang_code,
        "_detected_language_confidence": round(lang_conf, 3),
        "_track": track,
        "_role_category": role_category,
    })
    if len(description) < 200:
        raw_obj["_description_quality"] = "too_short"

    job_id = database.insert_job(
        title=title,
        company=company,
        location=location,
        job_board=source,
        url=url or None,
        description=description,
        posted_date=job.get("posted_date"),
        raw_data=json.dumps(raw_obj, ensure_ascii=False),
        dedup_hash=dedup_hash,
    )

    if job_id:
        database.update_job_classification(job_id, track, lang_code, role_category)
        stats["accepted"] = stats.get("accepted", 0) + 1
        stats.setdefault("by_track", {})[track] = stats.setdefault("by_track", {}).get(track, 0) + 1
        stats.setdefault("by_source", {})[source] = stats.setdefault("by_source", {}).get(source, 0) + 1
        return job_id

    if url:
        existing = database.get_job_by_url(url)
        if existing:
            stats["duplicates"] = stats.get("duplicates", 0) + 1
            return existing["id"]

    stats["errors"] = stats.get("errors", 0) + 1
    return None


# ── Watch mode ─────────────────────────────────────────────────────────────────

def _run_watch_mode(profile: dict, stats: dict) -> None:
    watchlist_path = config.BASE_DIR / "data" / "company_watchlist.json"
    if not watchlist_path.exists():
        print("[Watch] company_watchlist.json not found — aborting")
        return

    with open(watchlist_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    companies = data.get("companies", [])
    exclude_kws = profile.get("application_preferences", {}).get("exclude_title_keywords", [])

    print(f"[Watch] {len(companies)} companies in watchlist")

    for company in companies:
        name = company.get("name", "unknown")
        try:
            ats = company.get("ats_provider", "")
            ats_id = company.get("ats_id", "")

            raw_jobs = fetch_company(company)   # greenhouse/lever/ashby/smartrecruiters
            if not raw_jobs and ats not in _ATS_FETCHERS:
                continue

            if not raw_jobs:
                time.sleep(0.3)
                continue

            accepted_count = 0
            for job in raw_jobs:
                jid = _process_job(job, exclude_kws, stats)
                if jid and stats.get("accepted", 0) > (stats.get("_prev_accepted", 0)):
                    accepted_count += 1
                    stats["_prev_accepted"] = stats.get("accepted", 0)

            print(f"  {name}: {len(raw_jobs)} fetched → {accepted_count} new")
            database.log_search(
                query=f"watchlist:{name}",
                job_board=ats,
                jobs_found=len(raw_jobs),
                filters_used={"ats_id": ats_id},
            )

        except Exception as exc:
            print(f"  [ERROR] {name}: {exc}")
            stats["errors"] = stats.get("errors", 0) + 1
            continue

        time.sleep(0.5)


# ── XLSX export ────────────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_TRACK_A_FILL = PatternFill("solid", fgColor="DDEEFF")
_TRACK_B_FILL = PatternFill("solid", fgColor="EEFFDD")
_UNCLASS_FILL = PatternFill("solid", fgColor="F5F5F5")


def _export_xlsx(output_path: Path) -> None:
    with database.get_connection() as conn:
        rows = conn.execute("""
            SELECT j.id, j.track, j.title, j.company, j.location,
                   j.job_board, j.role_category, j.url, j.language,
                   j.posted_date, j.fetched_date
            FROM jobs j
            LEFT JOIN scores s ON s.job_id = j.id
            WHERE s.id IS NULL AND j.is_expired = 0
            ORDER BY j.fetched_date DESC, j.id DESC
        """).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Unscored Jobs"

    headers = [
        "ID", "Track", "Title", "Company", "Location",
        "Source", "Role Category", "URL", "Language",
        "Posted Date", "Fetched Date",
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 18

    for row in rows:
        ws.append(list(row))
        row_num = ws.max_row
        track = row[1] or ""
        fill = _TRACK_A_FILL if track == "A" else (_TRACK_B_FILL if track == "B" else _UNCLASS_FILL)
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_num, column=col_idx).fill = fill
        url_val = row[7]
        if url_val:
            url_cell = ws.cell(row=row_num, column=8)
            url_cell.hyperlink = url_val
            url_cell.font = Font(color="0563C1", underline="single")
            url_cell.value = "Link"

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    col_widths = [6, 8, 40, 25, 22, 14, 22, 8, 10, 14, 14]
    for i, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


# ── Summary ────────────────────────────────────────────────────────────────────

def _print_summary(stats: dict, xlsx_path: Optional[Path]) -> None:
    with database.get_connection() as conn:
        total_db = conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
        unscored = conn.execute(
            "SELECT COUNT(*) FROM jobs j LEFT JOIN scores s ON s.job_id = j.id WHERE s.id IS NULL"
        ).fetchone()[0]

    print("\n" + "═" * 54)
    print("  SEARCH AGENT — RUN SUMMARY")
    print("═" * 54)
    print(f"  Jobs fetched (all ATS)     : {stats.get('fetched', 0)}")
    print(f"  Jobs accepted (new)        : {stats.get('accepted', 0)}")
    print(f"  Duplicates skipped         : {stats.get('duplicates', 0)}")
    print(f"  Rejected — excluded title  : {stats.get('rejected_excluded', 0)}")
    print(f"  Rejected — seniority       : {stats.get('rejected_seniority', 0)}")
    for _reason, _n in sorted(stats.get("by_seniority_reason", {}).items()):
        print(f"    ↳ {_reason:<34}: {_n}")
    print(f"  Rejected — geography       : {stats.get('rejected_geography', 0)}")
    print(f"  Rejected — language        : {stats.get('rejected_language', 0)}")
    print(f"  Rejected — missing fields  : {stats.get('rejected_other', 0)}")
    print("─" * 54)
    for track, n in sorted(stats.get("by_track", {}).items()):
        print(f"  New Track {track:<13}        : {n}")
    for source, n in sorted(stats.get("by_source", {}).items()):
        print(f"  New via {source:<16}  : {n}")
    print("─" * 54)
    print(f"  Total jobs in database     : {total_db}")
    print(f"  Jobs awaiting scoring      : {unscored}")
    if xlsx_path and xlsx_path.exists():
        print(f"  Export                     : {xlsx_path.name}")
    print("═" * 54)


# ── Entry point ────────────────────────────────────────────────────────────────

def _run_public_boards(profile: dict, stats: dict) -> None:
    """Search aggregators/portals (Adzuna free API + LinkedIn/Indeed/WTTJ via Apify)
    and process results through the same dedup/gating/persistence as ATS jobs.
    No-op for any source whose credentials are absent."""
    from agents import adzuna_search, apify_search

    jobs = adzuna_search.search_adzuna(profile) + apify_search.search_public_boards(profile)
    if not jobs:
        return
    print(f"[Public] {len(jobs)} job(s) from public boards (LinkedIn/Indeed/WTTJ)")
    exclude_kws = profile.get("application_preferences", {}).get("exclude_title_keywords", [])
    for job in jobs:
        _process_job(job, exclude_kws, stats)
    database.log_search(query="public_boards", job_board="apify",
                        jobs_found=len(jobs), filters_used={})


def _run_career_sites(profile: dict, stats: dict) -> None:
    """Search company career pages (JSON-LD / RSS) for jobs from companies on no
    ATS. Processed through the same dedup/gating/persistence. No-op if no config."""
    from agents import career_site_search

    jobs = career_site_search.search_career_sites()
    if not jobs:
        return
    print(f"[CareerSites] {len(jobs)} job(s) from company career pages")
    exclude_kws = profile.get("application_preferences", {}).get("exclude_title_keywords", [])
    for job in jobs:
        _process_job(job, exclude_kws, stats)
    database.log_search(query="career_sites", job_board="career_site",
                        jobs_found=len(jobs), filters_used={})


def run() -> None:
    profile = load_profile()
    stats: dict = {}

    print("\n[Search Agent] ATS watchlist run starting...")
    _run_watch_mode(profile, stats)
    _run_career_sites(profile, stats)    # company career pages (JSON-LD / RSS)
    _run_public_boards(profile, stats)   # LinkedIn/Indeed/WTTJ via Apify (if enabled)

    today_str = date.today().isoformat()
    xlsx_path = Path(config.OUTPUTS_DIR) / "exports" / f"jobs_master_{today_str}.xlsx"

    if stats.get("fetched", 0) > 0:
        _export_xlsx(xlsx_path)
    else:
        xlsx_path = None

    _print_summary(stats, xlsx_path)


if __name__ == "__main__":
    from dotenv import load_dotenv
    load_dotenv()
    from core.database import initialize
    initialize()
    run()
