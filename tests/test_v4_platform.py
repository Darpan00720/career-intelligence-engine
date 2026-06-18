"""Tests for the v4 production platform:

  agent lifecycle, orchestrator, scheduler, REST API, review queue,
  notifications, feedback engine, experiments, parallel execution, and the
  system-health dashboard.
"""
import sqlite3
import tempfile
import time
import unittest
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

import openpyxl

from agents.base_agent import BaseAgent
from core import (
    concurrency,
    database,
    feedback_engine,
    notifications,
    orchestrator,
    review_queue,
    scheduler,
    system_health,
)


# ── Shared in-memory DB ───────────────────────────────────────────────────────

def _mem_conn():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        CREATE TABLE jobs (id INTEGER PRIMARY KEY, title TEXT, company TEXT,
            location TEXT, role_category TEXT, visa_gate TEXT, work_mode TEXT,
            url TEXT, fetched_date DATE);
        CREATE TABLE scores (id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER,
            total_score INTEGER);
        CREATE TABLE applications (id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER,
            status TEXT, outcome TEXT, notes TEXT, last_updated DATETIME);
        CREATE TABLE company_research (id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER);
        CREATE TABLE documents (id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER,
            type TEXT, version INTEGER);
        CREATE TABLE review_queue (id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER,
            state TEXT DEFAULT 'Pending Review', recommendation TEXT,
            resume_doc_id INTEGER, cover_letter_doc_id INTEGER, notes TEXT,
            created_at DATETIME DEFAULT (DATETIME('now')),
            updated_at DATETIME DEFAULT (DATETIME('now')));
        CREATE TABLE experiments (id INTEGER PRIMARY KEY AUTOINCREMENT,
            experiment_id TEXT UNIQUE, name TEXT, status TEXT DEFAULT 'running',
            variants TEXT, winner TEXT, created_at DATETIME);
        CREATE TABLE experiment_events (id INTEGER PRIMARY KEY AUTOINCREMENT,
            experiment_id TEXT, variant TEXT, metric TEXT, value REAL,
            created_at DATETIME);
        CREATE TABLE pipeline_runs (id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id TEXT UNIQUE, duration_seconds REAL, detail TEXT);
    """)
    return conn


class MemDBTest(unittest.TestCase):
    def setUp(self):
        self.conn = _mem_conn()

        @contextmanager
        def _ctx():
            yield self.conn
            self.conn.commit()

        self._p = patch("core.database.get_connection", _ctx)
        self._p.start()

    def tearDown(self):
        self._p.stop()
        self.conn.close()

    def _seed_job(self, jid, company="Acme", score=90, status="Applied",
                  role="ai_strategy"):
        self.conn.execute(
            "INSERT INTO jobs (id,title,company,role_category,visa_gate,fetched_date) "
            "VALUES (?,?,?,?,?,?)", (jid, f"Role {jid}", company, role, "PASS", "2026-06-10"))
        self.conn.execute("INSERT INTO scores (job_id,total_score) VALUES (?,?)", (jid, score))
        self.conn.execute("INSERT INTO applications (job_id,status) VALUES (?,?)", (jid, status))
        self.conn.commit()


# ── 1. Agent lifecycle ────────────────────────────────────────────────────────

class _LifecycleAgent(BaseAgent):
    name = "lifecycle"

    def __init__(self, calls, fail=False, valid=True, **kw):
        super().__init__(**kw)
        self.calls, self._fail, self._valid = calls, fail, valid

    def initialize(self): self.calls.append("initialize")
    def validate(self): self.calls.append("validate"); return self._valid
    def run(self):
        self.calls.append("run")
        if self._fail:
            raise RuntimeError("boom")
        return {"done": True}
    def cleanup(self): self.calls.append("cleanup")


class TestAgentLifecycle(unittest.TestCase):
    def test_hooks_run_in_order(self):
        calls = []
        res = _LifecycleAgent(calls).execute()
        self.assertEqual(calls, ["initialize", "validate", "run", "cleanup"])
        self.assertTrue(res.ok)
        self.assertEqual(res.data, {"done": True})

    def test_validation_failure_skips_run(self):
        calls = []
        res = _LifecycleAgent(calls, valid=False).execute()
        self.assertEqual(res.status, "skipped")
        self.assertNotIn("run", calls)
        self.assertIn("cleanup", calls)   # cleanup still runs

    def test_run_error_captured_and_cleanup_called(self):
        calls = []
        res = _LifecycleAgent(calls, fail=True).execute()
        self.assertEqual(res.status, "error")
        self.assertIn("boom", res.error)
        self.assertEqual(calls[-1], "cleanup")

    def test_dependency_injection(self):
        class A(BaseAgent):
            name = "di"
            def run(self): return {"v": self.dep("fn")()}
        res = A(deps={"fn": lambda: 42}).execute()
        self.assertEqual(res.data["v"], 42)


# ── 2. Orchestrator ───────────────────────────────────────────────────────────

class _StubAgent(BaseAgent):
    def __init__(self, name, log, fail_times=0):
        super().__init__()
        self.name, self.log, self.fail_times = name, log, fail_times
        self.attempts = 0
    def run(self):
        self.attempts += 1
        self.log.append(self.name)
        if self.attempts <= self.fail_times:
            raise RuntimeError("transient")
        return {"agent": self.name}


class TestOrchestrator(unittest.TestCase):
    def test_sequential_order(self):
        log = []
        orch = orchestrator.Orchestrator()
        for n in ("a", "b", "c"):
            orch.register(_StubAgent(n, log))
        report = orch.run([["a"], ["b"], ["c"]], mode="sequential")
        self.assertEqual(log, ["a", "b", "c"])
        self.assertTrue(report.succeeded)
        self.assertEqual(len(report.results), 3)

    def test_retry_on_failure(self):
        log = []
        orch = orchestrator.Orchestrator(max_retries=2)
        orch.register(_StubAgent("flaky", log, fail_times=1))
        report = orch.run([["flaky"]])
        self.assertTrue(report.succeeded)        # succeeded on retry
        self.assertEqual(log.count("flaky"), 2)  # one fail + one success

    def test_failure_recorded_after_retries_exhausted(self):
        log = []
        orch = orchestrator.Orchestrator(max_retries=1)
        orch.register(_StubAgent("bad", log, fail_times=5))
        report = orch.run([["bad"]])
        self.assertFalse(report.succeeded)
        self.assertIn("bad", report.failures)

    def test_parallel_stage_runs_all(self):
        log = []
        orch = orchestrator.Orchestrator()
        for n in ("x", "y", "z"):
            orch.register(_StubAgent(n, log))
        report = orch.run([["x", "y", "z"]], mode="parallel")
        self.assertEqual(sorted(log), ["x", "y", "z"])
        self.assertEqual(len(report.results), 3)


# ── 3. Scheduler ──────────────────────────────────────────────────────────────

class TestScheduler(unittest.TestCase):
    def test_cron_matches(self):
        self.assertTrue(scheduler.cron_matches("0 8 * * *", datetime(2026, 6, 18, 8, 0)))
        self.assertFalse(scheduler.cron_matches("0 8 * * *", datetime(2026, 6, 18, 9, 0)))

    def test_daily_fires_once_per_minute(self):
        hits = []
        sc = scheduler.Scheduler()
        sc.add("nightly", scheduler.daily(8, 0), lambda: hits.append(1))
        sc.tick(datetime(2026, 6, 18, 8, 0))
        sc.tick(datetime(2026, 6, 18, 8, 0))   # same minute → no double fire
        sc.tick(datetime(2026, 6, 18, 9, 0))   # not due
        self.assertEqual(len(hits), 1)

    def test_pause_resume_disable(self):
        hits = []
        sc = scheduler.Scheduler()
        sc.add("j", scheduler.daily(8), lambda: hits.append(1))
        sc.pause("j"); sc.tick(datetime(2026, 6, 18, 8, 0)); self.assertEqual(hits, [])
        sc.resume("j"); sc.tick(datetime(2026, 6, 18, 8, 0)); self.assertEqual(len(hits), 1)
        sc.disable("j"); sc.tick(datetime(2026, 6, 19, 8, 0)); self.assertEqual(len(hits), 1)

    def test_weekly_and_history_with_retry(self):
        sc = scheduler.Scheduler()
        sc.add("w", scheduler.weekly(0, 9), lambda: (_ for _ in ()).throw(RuntimeError("x")),
               max_retries=2)
        # 2026-06-21 is a Sunday (dow 0)
        sc.tick(datetime(2026, 6, 21, 9, 0))
        hist = sc.history("w")
        self.assertEqual(len(hist), 1)
        self.assertEqual(hist[0]["status"], "error")
        self.assertEqual(hist[0]["attempts"], 3)   # 1 + 2 retries


# ── 4. REST API ───────────────────────────────────────────────────────────────

_API_ROWS = [
    {"id": 1, "company": "Acme", "title": "AI PM", "location": "Milan",
     "total_score": 95, "fetched_date": "2026-06-10", "role_category": "ai_strategy",
     "visa_gate": "PASS", "url": "http://a/1", "application_status": "Not Started"},
    {"id": 2, "company": "Beta", "title": "Strategy", "location": "London",
     "total_score": 72, "fetched_date": "2026-06-11", "role_category": "business_strategy",
     "visa_gate": "PASS", "url": "http://b/2", "application_status": "Applied"},
]


class TestRestApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        from api.app import app
        cls.client = TestClient(app)

    def test_jobs_pagination_and_filter(self):
        with patch("core.database.get_all_scored_jobs_ranked", return_value=_API_ROWS):
            r = self.client.get("/api/jobs?limit=1&offset=0&min_score=80")
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual(body["total"], 1)      # only the 95-score job passes min_score
        self.assertEqual(body["limit"], 1)

    def test_jobs_sort_validation(self):
        with patch("core.database.get_all_scored_jobs_ranked", return_value=_API_ROWS):
            r = self.client.get("/api/jobs?sort=nonsense")
        self.assertEqual(r.status_code, 400)

    def test_top_and_recommendations(self):
        with patch("core.database.get_all_scored_jobs_ranked", return_value=_API_ROWS), \
             patch("core.database.get_researched_job_ids", return_value=set()), \
             patch("core.profile_loader.load", return_value={}):
            top = self.client.get("/api/jobs/top?min_score=70")
            recs = self.client.get("/api/recommendations?min_score=70")
        self.assertEqual(top.status_code, 200)
        self.assertEqual(recs.status_code, 200)
        self.assertTrue(all("recommendation" in i for i in recs.json()["items"]))

    def test_update_status(self):
        with patch("core.tracker.set_status") as setter:
            r = self.client.post("/api/applications/status",
                                 json={"job_id": 1, "status": "Applied"})
        self.assertEqual(r.status_code, 200)
        setter.assert_called_once()

    def test_update_status_invalid(self):
        r = self.client.post("/api/applications/status",
                             json={"job_id": 1, "status": "Ghosted"})
        self.assertEqual(r.status_code, 400)

    def test_pipeline_run_endpoint(self):
        with patch("core.pipeline.run_autonomous_pipeline",
                   return_value={"run_id": "abc", "status": "SUCCESS", "duration_seconds": 1.0}):
            r = self.client.post("/api/pipeline/run", json={"top_n": 5})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["data"]["status"], "SUCCESS")


# ── 5. Review queue ───────────────────────────────────────────────────────────

class TestReviewQueue(MemDBTest):
    def test_enqueue_and_transitions(self):
        self._seed_job(1)
        rid = review_queue.enqueue(1, recommendation="Apply Now")
        self.assertEqual(review_queue.get(rid)["state"], review_queue.PENDING)
        self.assertEqual(len(review_queue.pending()), 1)

        review_queue.approve(rid, notes="looks good")
        self.assertEqual(review_queue.get(rid)["state"], review_queue.APPROVED)
        self.assertEqual(review_queue.get(rid)["notes"], "looks good")
        self.assertEqual(len(review_queue.pending()), 0)

    def test_reject_and_regenerate(self):
        self._seed_job(1)
        rid = review_queue.enqueue(1)
        review_queue.request_regeneration(rid)
        self.assertEqual(review_queue.get(rid)["state"], review_queue.REGENERATE)
        review_queue.reject(rid)
        self.assertEqual(review_queue.get(rid)["state"], review_queue.REJECTED)

    def test_invalid_state(self):
        self._seed_job(1)
        rid = review_queue.enqueue(1)
        with self.assertRaises(review_queue.InvalidReviewState):
            review_queue._transition(rid, "Bogus")


# ── 6. Notifications ──────────────────────────────────────────────────────────

class _RecordingChannel(notifications.Channel):
    name = "recording"
    def __init__(self): self.sent = []
    def send(self, subject, message, level="INFO"):
        self.sent.append((subject, message, level)); return True


class TestNotifications(unittest.TestCase):
    def test_dispatch_to_channels(self):
        ch = _RecordingChannel()
        n = notifications.Notifier(channels=[ch])
        delivered = n.dispatch("offer_received", "Offer!", level="INFO")
        self.assertEqual(delivered, {"recording": True})
        self.assertEqual(len(ch.sent), 1)

    def test_unconfigured_channels_noop(self):
        with patch.dict("os.environ", {}, clear=True):
            self.assertFalse(notifications.EmailChannel().send("s", "m"))
            self.assertFalse(notifications.SlackChannel().send("s", "m"))

    def test_notification_agent(self):
        from agents.career_agents import NotificationAgent
        ch = _RecordingChannel()
        agent = NotificationAgent(
            context={"events": [notifications.offer_event("Acme", "AI PM"),
                                notifications.high_scoring_job_event("Beta", "PM", 95)]},
            deps={"notifier": notifications.Notifier(channels=[ch])},
        )
        res = agent.execute()
        self.assertTrue(res.ok)
        self.assertEqual(res.data["notifications_sent"], 2)
        self.assertEqual(len(ch.sent), 2)


# ── 7. Feedback engine ────────────────────────────────────────────────────────

class TestFeedbackEngine(MemDBTest):
    def test_conversion_by_company_and_band(self):
        self._seed_job(1, company="Acme", score=95, status="Offer")
        self._seed_job(2, company="Acme", score=92, status="Interview")
        self._seed_job(3, company="Beta", score=72, status="Applied")
        by_co = feedback_engine.conversion_by_company()
        self.assertEqual(by_co["Acme"]["applied"], 2)
        self.assertEqual(by_co["Acme"]["interviews"], 2)
        self.assertEqual(by_co["Acme"]["offers"], 1)
        self.assertEqual(by_co["Acme"]["offer_rate"], 50.0)
        bands = feedback_engine.conversion_by_score_band()
        self.assertEqual(bands["90-100"]["interviews"], 2)
        self.assertEqual(bands["70-79"]["interviews"], 0)

    def test_recommendation_effectiveness_runs(self):
        self._seed_job(1, score=95, status="Offer")
        eff = feedback_engine.recommendation_effectiveness()
        self.assertTrue(eff)   # produces at least one recommendation bucket


# ── 8. Experiments ────────────────────────────────────────────────────────────

class TestExperiments(MemDBTest):
    def test_create_assign_record_winner(self):
        import experiments
        experiments.create_experiment("exp1", ["control", "treatment"], name="Test")
        v1 = experiments.assign_variant("exp1", "unit-A")
        v2 = experiments.assign_variant("exp1", "unit-A")
        self.assertEqual(v1, v2)                     # deterministic
        self.assertIn(v1, ["control", "treatment"])

        experiments.record_metric("exp1", "control", "interview_rate", 0.2)
        experiments.record_metric("exp1", "treatment", "interview_rate", 0.6)
        res = experiments.results("exp1", "interview_rate")
        self.assertEqual(res["treatment"]["interview_rate"]["mean"], 0.6)
        winner = experiments.decide_winner("exp1", "interview_rate")
        self.assertEqual(winner, "treatment")
        self.assertEqual(database.get_experiment("exp1")["winner"], "treatment")

    def test_requires_two_variants(self):
        import experiments
        with self.assertRaises(ValueError):
            experiments.create_experiment("bad", ["only"])


# ── 9. Parallel execution ─────────────────────────────────────────────────────

class TestConcurrency(unittest.TestCase):
    def test_parallel_map_preserves_order(self):
        res = concurrency.parallel_map(lambda x: x * 2, [1, 2, 3, 4], max_workers=3)
        self.assertEqual([r.value for r in res], [2, 4, 6, 8])
        self.assertTrue(all(r.ok for r in res))

    def test_parallel_map_captures_errors(self):
        def f(x):
            if x == 2:
                raise ValueError("nope")
            return x
        res = concurrency.parallel_map(f, [1, 2, 3], max_workers=2)
        self.assertTrue(res[0].ok)
        self.assertFalse(res[1].ok)
        self.assertIn("nope", res[1].error)

    def test_parallel_runs_concurrently(self):
        # 4 tasks of 0.1s across 4 workers should finish well under 0.4s.
        start = time.perf_counter()
        concurrency.parallel_map(lambda _: time.sleep(0.1), range(4), max_workers=4)
        self.assertLess(time.perf_counter() - start, 0.35)

    def test_backoff_retry_succeeds_after_failures(self):
        calls = {"n": 0}
        def flaky():
            calls["n"] += 1
            if calls["n"] < 3:
                raise RuntimeError("x")
            return "ok"
        self.assertEqual(concurrency.backoff_retry(flaky, retries=3, base_delay=0.0), "ok")
        self.assertEqual(calls["n"], 3)

    def test_rate_limiter(self):
        rl = concurrency.RateLimiter(min_interval=0.05)
        start = time.perf_counter()
        rl.wait(); rl.wait()
        self.assertGreaterEqual(time.perf_counter() - start, 0.05)


# ── 10. System health dashboard ───────────────────────────────────────────────

class TestSystemHealth(unittest.TestCase):
    def test_export_creates_workbook_with_charts(self):
        with tempfile.TemporaryDirectory() as td:
            out = str(Path(td) / "system_health.xlsx")
            with patch("core.analytics.application_metrics", return_value={
                        "jobs_found": 100, "jobs_scored": 80, "researched": 10,
                        "documents": 5, "applied": 4, "assessments": 2,
                        "interviews": 2, "final_rounds": 1, "offers": 1, "rejections": 1}), \
                 patch("core.analytics.conversion_metrics", return_value={
                        "conversion_rate": 5.0, "interview_rate": 50.0, "offer_rate": 25.0}), \
                 patch("core.analytics.funnel_metrics", return_value=[
                        {"stage": "Jobs Found", "count": 100}, {"stage": "Applied", "count": 4}]), \
                 patch("core.database.get_pipeline_runs", return_value=[
                        {"duration_seconds": 10.0,
                         "detail": '{"errors": [], "research_researched": 8, '
                                   '"research_cached": 2, "documents_generated": 3, '
                                   '"documents_reused": 7}'}]), \
                 patch("core.database.get_document_counts", return_value={"total": 10, "reused_versions": 4}), \
                 patch("core.database.get_average_score", return_value=68.5), \
                 patch("core.database.get_jobs_per_day", return_value=[
                        {"day": "2026-06-10", "count": 30}, {"day": "2026-06-11", "count": 25}]):
                metrics = system_health.collect_metrics()
                path = system_health.export_system_health_xlsx(out)

                self.assertEqual(metrics["research_cache_hit_rate"], 20.0)   # 2 / (8+2)
                self.assertEqual(metrics["document_reuse_rate"], 70.0)       # 7 / (3+7)
                wb = openpyxl.load_workbook(path)
                self.assertIn("Health", wb.sheetnames)
                self.assertIn("Funnel", wb.sheetnames)
                self.assertGreaterEqual(
                    len(wb["Health"]._charts) + len(wb["Funnel"]._charts), 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
