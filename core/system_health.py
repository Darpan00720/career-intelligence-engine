"""System Health Dashboard (v4) → outputs/system_health.xlsx.

Aggregates operational + outcome metrics from the database and persisted
pipeline runs, with charts. Cache-hit and document-reuse rates are derived from
the stats each pipeline run records in pipeline_runs.detail, so they reflect
real cumulative behavior across runs.
"""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill

from core import analytics, config, database

_HEADER_FILL = PatternFill(start_color="FF263238", end_color="FF263238", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_TITLE_FONT = Font(bold=True, size=14)


def _run_derived_rates() -> dict:
    """Aggregate cache-hit and document-reuse rates across persisted runs."""
    researched = cached = generated = reused = 0
    durations: list[float] = []
    failures = 0
    for run in database.get_pipeline_runs(limit=200):
        durations.append(run.get("duration_seconds") or 0.0)
        try:
            detail = json.loads(run.get("detail") or "{}")
        except (json.JSONDecodeError, TypeError):
            detail = {}
        researched += detail.get("research_researched", 0)
        cached += detail.get("research_cached", 0)
        generated += detail.get("documents_generated", 0)
        reused += detail.get("documents_reused", 0)
        failures += len(detail.get("errors", []) or [])
    cache_total = researched + cached
    doc_total = generated + reused
    avg_dur = round(sum(durations) / len(durations), 1) if durations else 0.0
    return {
        "research_cache_hit_rate": round(100.0 * cached / cache_total, 1) if cache_total else 0.0,
        "document_reuse_rate": round(100.0 * reused / doc_total, 1) if doc_total else 0.0,
        "avg_pipeline_duration": avg_dur,
        "agent_failures": failures,
        "runs": len(durations),
    }


def collect_metrics() -> dict:
    """Gather all system-health metrics into a single dict."""
    am = analytics.application_metrics()
    cm = analytics.conversion_metrics()
    derived = _run_derived_rates()
    docs = database.get_document_counts()
    return {
        "jobs_total": am["jobs_found"],
        "average_score": database.get_average_score(),
        "jobs_per_day": database.get_jobs_per_day(limit_days=30),
        "research_cache_hit_rate": derived["research_cache_hit_rate"],
        "document_reuse_rate": derived["document_reuse_rate"],
        "documents_total": docs["total"],
        "avg_pipeline_duration": derived["avg_pipeline_duration"],
        "agent_failures": derived["agent_failures"],
        "pipeline_runs": derived["runs"],
        "conversion_rate": cm["conversion_rate"],
        "interview_rate": cm["interview_rate"],
        "offer_rate": cm["offer_rate"],
    }


def export_system_health_xlsx(output_path: str | None = None) -> str:
    """Write the system-health workbook (Health sheet + Discovery chart)."""
    m = collect_metrics()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Health"

    ws.append(["System Health"])
    ws["A1"].font = _TITLE_FONT
    ws.append([])

    ws.append(["Metric", "Value"])
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    cards = [
        ("Jobs Discovered (total)", m["jobs_total"]),
        ("Average Score", m["average_score"]),
        ("Research Cache Hit Rate %", m["research_cache_hit_rate"]),
        ("Document Reuse Rate %", m["document_reuse_rate"]),
        ("Documents Generated (total)", m["documents_total"]),
        ("Avg Pipeline Duration (s)", m["avg_pipeline_duration"]),
        ("Agent Failures", m["agent_failures"]),
        ("Pipeline Runs", m["pipeline_runs"]),
        ("Conversion Rate %", m["conversion_rate"]),
        ("Interview Rate %", m["interview_rate"]),
        ("Offer Rate %", m["offer_rate"]),
    ]
    for label, value in cards:
        ws.append([label, value])

    # Jobs discovered/day table + line chart.
    ws.append([])
    chart_header_row = ws.max_row + 1
    ws.append(["Day", "Jobs Discovered"])
    for cell in ws[chart_header_row]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    per_day = list(reversed(m["jobs_per_day"]))  # chronological
    for item in per_day:
        ws.append([item["day"], item["count"]])
    last = ws.max_row

    if last > chart_header_row:
        chart = LineChart()
        chart.title = "Jobs Discovered / Day"
        data = Reference(ws, min_col=2, min_row=chart_header_row, max_row=last)
        cats = Reference(ws, min_col=1, min_row=chart_header_row + 1, max_row=last)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height, chart.width = 8, 18
        ws.add_chart(chart, "D3")

    # Conversion funnel bar chart on a second sheet.
    ws2 = wb.create_sheet("Funnel")
    ws2.append(["Stage", "Count"])
    for cell in ws2[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
    for item in analytics.funnel_metrics():
        ws2.append([item["stage"], item["count"]])
    bar = BarChart()
    bar.title = "Application Funnel"
    data = Reference(ws2, min_col=2, min_row=1, max_row=ws2.max_row)
    cats = Reference(ws2, min_col=1, min_row=2, max_row=ws2.max_row)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    ws2.add_chart(bar, "D2")

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16

    out = Path(output_path) if output_path else Path(config.OUTPUTS_DIR) / "system_health.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out.resolve())
