"""
Tests for the autonomous ranking + jobs_master export pipeline.

Covers:
  - ranking logic (order, 1-based ranks, score-then-date tiebreak, no mutation)
  - application-priority band assignment (90/80/70 thresholds)
  - jobs_master Excel export ordering + exact column layout + derived fields
  - autonomous pipeline execution (step orchestration + console progress)
  - backward compatibility (existing strategic buckets untouched)
"""
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from core import database, exporter
from core.ranking import rank_jobs, top_jobs
from core.scorer import (
    APPLICATION_PRIORITY_THRESHOLDS,
    application_priority,
    assign_priority_bucket,
)


# ── Ranking logic ───────────────────────────────────────────────────────────────

class TestRankJobs(unittest.TestCase):
    def test_orders_by_score_desc_and_assigns_ranks(self):
        rows = [
            {"id": 1, "total_score": 84, "fetched_date": "2026-06-01"},
            {"id": 2, "total_score": 96, "fetched_date": "2026-06-01"},
            {"id": 3, "total_score": 88, "fetched_date": "2026-06-01"},
            {"id": 4, "total_score": 92, "fetched_date": "2026-06-01"},
        ]
        ranked = rank_jobs(rows)
        self.assertEqual([r["rank"] for r in ranked], [1, 2, 3, 4])
        self.assertEqual([r["total_score"] for r in ranked], [96, 92, 88, 84])
        self.assertEqual([r["id"] for r in ranked], [2, 4, 3, 1])

    def test_tiebreak_by_date_found_desc(self):
        rows = [
            {"id": 1, "total_score": 80, "fetched_date": "2026-06-01"},
            {"id": 2, "total_score": 80, "fetched_date": "2026-06-10"},
        ]
        ranked = rank_jobs(rows)
        # Equal scores → most recent first
        self.assertEqual(ranked[0]["id"], 2)
        self.assertEqual(ranked[1]["id"], 1)

    def test_does_not_mutate_input(self):
        rows = [{"id": 1, "total_score": 50, "fetched_date": "2026-06-01"}]
        rank_jobs(rows)
        self.assertNotIn("rank", rows[0])
        self.assertNotIn("application_priority", rows[0])

    def test_handles_missing_and_none_scores(self):
        rows = [
            {"id": 1, "total_score": None, "fetched_date": "2026-06-01"},
            {"id": 2, "fetched_date": "2026-06-01"},
            {"id": 3, "total_score": 70, "fetched_date": "2026-06-01"},
        ]
        ranked = rank_jobs(rows)
        self.assertEqual(ranked[0]["id"], 3)  # only real score ranks first
        self.assertEqual(len(ranked), 3)

    def test_empty_input(self):
        self.assertEqual(rank_jobs([]), [])


# ── Application priority bands ───────────────────────────────────────────────────

class TestApplicationPriority(unittest.TestCase):
    def test_thresholds(self):
        self.assertEqual(application_priority(96), "Apply Now")
        self.assertEqual(application_priority(90), "Apply Now")
        self.assertEqual(application_priority(89), "This Week")
        self.assertEqual(application_priority(80), "This Week")
        self.assertEqual(application_priority(79), "Good Match")
        self.assertEqual(application_priority(70), "Good Match")
        self.assertEqual(application_priority(69), "Low Priority")
        self.assertEqual(application_priority(0), "Low Priority")

    def test_none_score_is_low_priority(self):
        self.assertEqual(application_priority(None), "Low Priority")

    def test_thresholds_table_covers_full_range(self):
        # Every score 0..100 maps to exactly one band.
        for s in range(0, 101):
            self.assertIn(
                application_priority(s),
                {label for _, label in APPLICATION_PRIORITY_THRESHOLDS},
            )

    def test_strategic_bucket_taxonomy_unchanged(self):
        # Backward compatibility: the strategic buckets are independent and intact.
        self.assertEqual(assign_priority_bucket(90), "Apply Immediately")
        self.assertEqual(assign_priority_bucket(70), "High Priority")
        self.assertEqual(assign_priority_bucket(0), "Ignore")


# ── Derived export fields ────────────────────────────────────────────────────────

class TestDerivedFields(unittest.TestCase):
    def test_country(self):
        self.assertEqual(exporter._derive_country("Milan, Italy"), "Italy")
        self.assertEqual(exporter._derive_country("London"), "United Kingdom")
        self.assertEqual(exporter._derive_country("Remote - Europe"), "Remote/EU")
        self.assertEqual(exporter._derive_country(""), "Unknown")

    def test_work_mode(self):
        self.assertEqual(exporter._derive_work_mode("Hybrid - Berlin", ""), "Hybrid")
        self.assertEqual(exporter._derive_work_mode("Remote", ""), "Remote")
        self.assertEqual(exporter._derive_work_mode("Milan", "office based role"), "On-site")

    def test_sponsorship(self):
        self.assertEqual(exporter._derive_sponsorship({"visa_gate": "PASS"}), "Likely")
        self.assertEqual(exporter._derive_sponsorship({"visa_gate": "FAIL"}), "No")
        self.assertEqual(exporter._derive_sponsorship({}), "Unknown")


# ── jobs_master Excel export ─────────────────────────────────────────────────────

_EXPECTED_HEADERS = [
    "Rank", "Score", "Priority", "Company", "Job Title", "Location", "Country",
    "Work Mode", "Sponsorship", "Category", "Application Status", "Date Found",
    "Job URL", "Notes",
]


def _sample_rows():
    return [
        {"id": 1, "title": "AI Product Intern", "company": "Acme", "location": "Milan, Italy",
         "description": "remote ai product", "role_category": "ai_strategy", "url": "http://a/1",
         "fetched_date": "2026-06-10", "visa_gate": "PASS", "total_score": 96,
         "priority_bucket": "Apply Immediately", "application_status": None},
        {"id": 2, "title": "Strategy Graduate", "company": "Beta", "location": "Amsterdam, Netherlands",
         "description": "hybrid strategy", "role_category": "business_strategy", "url": "http://b/2",
         "fetched_date": "2026-06-11", "visa_gate": "FAIL", "total_score": 72,
         "priority_bucket": "High Priority", "application_status": "Applied"},
        {"id": 3, "title": "HR Analyst", "company": "Gamma", "location": "Rome, Italy",
         "description": "office", "role_category": "hr_analytics", "url": "http://c/3",
         "fetched_date": "2026-06-09", "visa_gate": None, "total_score": 40,
         "priority_bucket": "Low Priority", "application_status": None},
    ]


class TestMasterExport(unittest.TestCase):
    def _export(self, rows):
        td = tempfile.mkdtemp()
        out = os.path.join(td, "jobs_master.xlsx")
        with patch.object(database, "get_all_scored_jobs_ranked", return_value=rows):
            path = exporter.export_jobs_master_xlsx(out)
        return path

    def test_exact_header_layout(self):
        path = self._export([])
        ws = openpyxl.load_workbook(path)["Jobs Master"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(_EXPECTED_HEADERS) + 1)]
        self.assertEqual(headers, _EXPECTED_HEADERS)

    def test_rows_sorted_by_score_desc(self):
        path = self._export(_sample_rows())
        ws = openpyxl.load_workbook(path)["Jobs Master"]
        scores = [ws.cell(row=r, column=2).value for r in range(2, 5)]
        self.assertEqual(scores, [96, 72, 40])
        ranks = [ws.cell(row=r, column=1).value for r in range(2, 5)]
        self.assertEqual(ranks, [1, 2, 3])

    def test_derived_and_priority_columns(self):
        path = self._export(_sample_rows())
        ws = openpyxl.load_workbook(path)["Jobs Master"]
        # Row 2 = top job (Acme, score 96)
        self.assertEqual(ws.cell(row=2, column=3).value, "Apply Now")   # Priority
        self.assertEqual(ws.cell(row=2, column=7).value, "Italy")       # Country
        self.assertEqual(ws.cell(row=2, column=9).value, "Likely")      # Sponsorship
        self.assertEqual(ws.cell(row=2, column=11).value, "Not Started")  # App status
        # Row 3 = Beta, score 72, applied
        self.assertEqual(ws.cell(row=3, column=3).value, "Good Match")
        self.assertEqual(ws.cell(row=3, column=11).value, "Applied")

    def test_notes_column_and_archive_copy(self):
        rows = _sample_rows()
        rows[0]["application_notes"] = "Referred by alumni"
        td = tempfile.mkdtemp()
        out = os.path.join(td, "jobs_master.xlsx")
        with patch.object(database, "get_all_scored_jobs_ranked", return_value=rows):
            exporter.export_jobs_master_xlsx(out)
        ws = openpyxl.load_workbook(out)["Jobs Master"]
        self.assertEqual(ws.cell(row=1, column=14).value, "Notes")
        self.assertEqual(ws.cell(row=2, column=14).value, "Referred by alumni")
        # A timestamped archive snapshot is written alongside the master file.
        archive = list(Path(td, "archive").glob("jobs_master_*.xlsx"))
        self.assertEqual(len(archive), 1)

    def test_default_path_is_jobs_master(self):
        from core import config
        with patch.object(database, "get_all_scored_jobs_ranked", return_value=[]):
            with tempfile.TemporaryDirectory() as td:
                with patch.object(config, "OUTPUTS_DIR", td):
                    path = exporter.export_jobs_master_xlsx()
                self.assertTrue(path.endswith("jobs_master.xlsx"))
                self.assertTrue(Path(path).exists())


# ── Integration: DB query + ranking + export against a real sqlite DB ───────────

def _make_db():
    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript("""
        CREATE TABLE jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT, company TEXT, location TEXT, url TEXT UNIQUE,
            description TEXT, role_category TEXT, visa_gate TEXT,
            fetched_date DATE, is_expired BOOLEAN DEFAULT 0
        );
        CREATE TABLE scores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER, total_score INTEGER, priority_bucket TEXT,
            explanation TEXT
        );
        CREATE TABLE applications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER, status TEXT, notes TEXT
        );
    """)
    jobs = [
        ("Lo", "C1", "Milan", "http://x/1", "d", "ai_strategy", "PASS", "2026-06-01", 55, "Medium Priority"),
        ("Hi", "C2", "London", "http://x/2", "d", "business_strategy", "PASS", "2026-06-02", 91, "Apply Immediately"),
        ("Mid", "C3", "Berlin", "http://x/3", "d", "hr_analytics", "FAIL", "2026-06-03", 75, "High Priority"),
    ]
    for j in jobs:
        cur = conn.execute(
            "INSERT INTO jobs (title,company,location,url,description,role_category,visa_gate,fetched_date) "
            "VALUES (?,?,?,?,?,?,?,?)", j[:8])
        conn.execute("INSERT INTO scores (job_id,total_score,priority_bucket,explanation) VALUES (?,?,?,?)",
                     (cur.lastrowid, j[8], j[9], "x"))
    conn.execute("INSERT INTO applications (job_id,status) VALUES (2,'Applied')")
    conn.commit()
    return conn


class TestRankingAgainstDb(unittest.TestCase):
    def test_get_all_scored_jobs_ranked_order_and_status(self):
        conn = _make_db()
        with patch("core.database.get_connection") as gc:
            gc.return_value.__enter__.return_value = conn
            rows = database.get_all_scored_jobs_ranked()
        self.assertEqual([r["total_score"] for r in rows], [91, 75, 55])
        ranked = rank_jobs(rows)
        self.assertEqual(ranked[0]["rank"], 1)
        self.assertEqual(ranked[0]["application_status"], "Applied")
        conn.close()

    def test_top_jobs_limit(self):
        conn = _make_db()
        with patch("core.database.get_connection") as gc:
            gc.return_value.__enter__.return_value = conn
            leaders = top_jobs(limit=2)
        self.assertEqual(len(leaders), 2)
        self.assertEqual([r["rank"] for r in leaders], [1, 2])
        conn.close()


# ── Autonomous pipeline orchestration ───────────────────────────────────────────

class TestAutonomousPipeline(unittest.TestCase):
    def test_runs_all_steps_in_order(self):
        from core import pipeline

        calls: list[str] = []

        with patch("agents.search_agent.run", side_effect=lambda: calls.append("search")), \
             patch("agents.scoring_agent.run", side_effect=lambda: calls.append("score")), \
             patch("agents.research_agent.run",
                   side_effect=lambda min_score=80: calls.append("research") or {"researched": 0}), \
             patch("agents.document_agent.run",
                   side_effect=lambda min_score=80: calls.append("documents") or {"resumes": 0}), \
             patch.object(pipeline, "_count_jobs", return_value=12), \
             patch.object(pipeline, "_count_scored", return_value=8), \
             patch("core.exporter.export_jobs_master_xlsx",
                   side_effect=lambda: calls.append("export") or "/out/jobs_master.xlsx"), \
             patch("core.database.get_all_scored_jobs_ranked", return_value=_sample_rows()), \
             patch("core.tracker.ensure_rows_for_scored",
                   side_effect=lambda: calls.append("tracker") or 3), \
             patch("core.database.insert_pipeline_run", return_value=1), \
             patch.object(pipeline, "display_recommendations"), \
             patch.object(pipeline, "display_top_jobs",
                          side_effect=lambda limit=20: calls.append("display")):
            summary = pipeline.run_autonomous_pipeline()

        self.assertEqual(
            calls,
            ["search", "score", "research", "documents", "export", "tracker", "display"],
        )
        self.assertEqual(summary["jobs_found"], 12)
        self.assertEqual(summary["jobs_scored"], 8)
        self.assertEqual(summary["jobs_ranked"], len(_sample_rows()))
        self.assertEqual(summary["tracker_rows_created"], 3)
        self.assertEqual(summary["export_path"], "/out/jobs_master.xlsx")
        self.assertEqual(summary["status"], "SUCCESS")
        self.assertIn("run_id", summary)

    def test_search_failure_does_not_abort_pipeline(self):
        from core import pipeline

        with patch("agents.search_agent.run", side_effect=RuntimeError("boom")), \
             patch("agents.scoring_agent.run"), \
             patch("agents.research_agent.run", return_value={}), \
             patch("agents.document_agent.run", return_value={}), \
             patch.object(pipeline, "_count_jobs", return_value=0), \
             patch.object(pipeline, "_count_scored", return_value=0), \
             patch("core.exporter.export_jobs_master_xlsx", return_value="/out/jobs_master.xlsx"), \
             patch("core.database.get_all_scored_jobs_ranked", return_value=[]), \
             patch("core.tracker.ensure_rows_for_scored", return_value=0), \
             patch("core.database.insert_pipeline_run", return_value=1), \
             patch.object(pipeline, "display_recommendations"), \
             patch.object(pipeline, "display_top_jobs"):
            summary = pipeline.run_autonomous_pipeline()

        self.assertEqual(summary["export_path"], "/out/jobs_master.xlsx")
        # Search failed but export succeeded → partial success, not aborted.
        self.assertEqual(summary["status"], "PARTIAL_SUCCESS")


class TestDisplayTopJobs(unittest.TestCase):
    def test_empty_message(self):
        from core import pipeline
        with patch("core.ranking.top_jobs", return_value=[]):
            rows = pipeline.display_top_jobs(limit=20)
        self.assertEqual(rows, [])

    def test_default_filters_below_70(self):
        from core import pipeline
        # _sample_rows scores: 96, 72, 40 → only 96 and 72 pass the >=70 default.
        with patch("core.ranking.top_jobs", return_value=rank_jobs(_sample_rows())):
            rows = pipeline.display_top_jobs(limit=20)
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r["total_score"] >= 70 for r in rows))
        # Status column is populated (defaulting when absent).
        self.assertTrue(all(r.get("application_status") for r in rows))

    def test_min_score_override_shows_all(self):
        from core import pipeline
        with patch("core.ranking.top_jobs", return_value=rank_jobs(_sample_rows())):
            rows = pipeline.display_top_jobs(limit=20, min_score=0)
        self.assertEqual(len(rows), 3)


if __name__ == "__main__":
    unittest.main(verbosity=2)
