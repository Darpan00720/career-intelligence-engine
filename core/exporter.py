"""
XLSX export for scored jobs (Phase 1.6b).

Exports all jobs that have a score row, color-coded by priority_bucket.
Requires openpyxl (already installed).

Usage:
    from core.exporter import export_jobs_xlsx
    path = export_jobs_xlsx()          # → outputs/exports/jobs_<timestamp>.xlsx
    path = export_jobs_xlsx("out.xlsx")  # explicit path
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core import config, database

# ── Bucket color palette ───────────────────────────────────────────────────────
# ARGB hex strings (no leading #). Alpha = FF (fully opaque).

_BUCKET_FILL: dict[str, PatternFill] = {
    b: PatternFill(start_color=c, end_color=c, fill_type="solid")
    for b, c in {
        "Apply Immediately": "FF00C853",   # vivid green
        "High Priority":     "FFB9F6CA",   # light green
        "Medium Priority":   "FFFFD740",   # amber
        "Low Priority":      "FFFFAB40",   # orange
        "Ignore":            "FFE0E0E0",   # mid-grey
        "Rejected":          "FFFFCDD2",   # light red/pink
    }.items()
}

_HEADER_FILL = PatternFill(start_color="FF263238", end_color="FF263238", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")

# ── Column definitions ────────────────────────────────────────────────────────
# (field_key, header_label, column_width)

_COLUMNS: list[tuple[str, str, int]] = [
    ("priority_bucket",          "Bucket",          18),
    ("total_score",              "Score",            7),
    ("eligibility_score",        "Elig. Score",      11),
    ("title",                    "Title",            38),
    ("company",                  "Company",          22),
    ("location",                 "Location",         20),
    ("posted_date",              "Posted",           12),
    ("role_category",            "Role Category",    22),
    ("track",                    "Track",            10),
    ("language_gate",            "Lang Gate",        11),
    ("visa_gate",                "Visa Gate",        10),
    ("detected_language",        "Detected Lang",    14),
    ("language_risk",            "Lang Risk",        11),
    ("eligibility_status",       "Elig. Status",     13),
    ("track_alignment_score",    "Track",             7),
    ("skill_match",              "Skills",            7),
    ("mba_relevance_score",      "MBA",               6),
    ("company_quality_score",    "Co. Quality",      11),
    ("intl_friendliness_score",  "Intl",              6),
    ("pivot_bonus_score",        "Pivot",             7),
    ("explanation",              "Explanation",      55),
    ("url",                      "URL",              45),
]


def export_jobs_xlsx(output_path: str | None = None) -> str:
    """
    Export all scored jobs to XLSX, color-coded by priority_bucket.

    Rows are ordered: Apply Immediately → High → Medium → Low → Ignore → Rejected.
    Within each bucket, rows are ordered by total_score descending.

    Returns the absolute path of the written file.
    """
    rows = database.get_scored_jobs_for_export()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scored Jobs"

    # ── Header row ─────────────────────────────────────────────────────────────
    headers = [col[1] for col in _COLUMNS]
    ws.append(headers)
    for col_idx, (_, _, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"

    # ── Data rows ──────────────────────────────────────────────────────────────
    for row_data in rows:
        values = [row_data.get(key) for key, _, _ in _COLUMNS]
        ws.append(values)

        bucket = row_data.get("priority_bucket") or "Ignore"
        fill   = _BUCKET_FILL.get(bucket)
        row_n  = ws.max_row

        for col_idx in range(1, len(_COLUMNS) + 1):
            cell = ws.cell(row=row_n, column=col_idx)
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=False)

        # Wrap explanation column only
        exp_col = next(
            i + 1 for i, (k, _, _) in enumerate(_COLUMNS) if k == "explanation"
        )
        ws.cell(row=row_n, column=exp_col).alignment = Alignment(
            vertical="top", wrap_text=True
        )

    # ── Summary sheet ──────────────────────────────────────────────────────────
    ws_sum = wb.create_sheet("Summary")
    _write_summary_sheet(ws_sum, rows)

    # ── Write file ─────────────────────────────────────────────────────────────
    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out = Path(config.OUTPUTS_DIR) / "exports" / f"jobs_{timestamp}.xlsx"
    else:
        out = Path(output_path)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out.resolve())


def _write_summary_sheet(ws, rows: list[dict]) -> None:
    """Append a compact bucket-count summary to the Summary worksheet."""
    from core.scorer import BUCKET_ORDER

    counts: dict[str, int] = {}
    for row in rows:
        b = row.get("priority_bucket") or "Ignore"
        counts[b] = counts.get(b, 0) + 1

    ws.append(["Bucket", "Count"])
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 8

    for bucket in BUCKET_ORDER:
        n = counts.get(bucket, 0)
        ws.append([bucket, n])
        fill = _BUCKET_FILL.get(bucket)
        if fill:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    ws.append([])
    ws.append(["Total", len(rows)])
    ws.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M")])


# ══════════════════════════════════════════════════════════════════════════════
#  jobs_master.xlsx — the prioritized application dashboard (single source of
#  truth). Distinct from export_jobs_xlsx above: a fixed, action-oriented column
#  set, ranked by score, color-coded by application priority band.
# ══════════════════════════════════════════════════════════════════════════════

# Application-priority color palette (matches core.scorer.application_priority).
_PRIORITY_FILL: dict[str, PatternFill] = {
    b: PatternFill(start_color=c, end_color=c, fill_type="solid")
    for b, c in {
        "Apply Now":    "FF00C853",   # vivid green
        "This Week":    "FFB9F6CA",   # light green
        "Good Match":   "FFFFD740",   # amber
        "Low Priority": "FFE0E0E0",   # mid-grey
    }.items()
}

# (field_key, header_label, column_width) — exact dashboard layout.
# column_width is the *minimum*; columns auto-fit wider content up to a cap.
_MASTER_COLUMNS: list[tuple[str, str, int]] = [
    ("rank",                 "Rank",               6),
    ("total_score",          "Score",              7),
    ("application_priority", "Priority",          13),
    ("company",              "Company",           24),
    ("title",                "Job Title",         42),
    ("location",             "Location",          22),
    ("country",              "Country",           12),
    ("work_mode",            "Work Mode",         11),
    ("sponsorship",          "Sponsorship",       12),
    ("role_category",        "Category",          22),
    ("application_status",   "Application Status", 18),
    ("fetched_date",         "Date Found",        12),
    ("url",                  "Job URL",           50),
    ("notes",                "Notes",             30),
]

# Upper bound for auto-fitted column width (keeps the sheet readable).
_MAX_COL_WIDTH = 60

# Location keyword → country label. First containment match wins; falls back to
# the raw location for remote/unknown values so nothing is silently dropped.
_COUNTRY_LOOKUP: list[tuple[str, str]] = [
    ("italy", "Italy"), ("italia", "Italy"), ("milan", "Italy"), ("milano", "Italy"),
    ("rome", "Italy"), ("roma", "Italy"), ("turin", "Italy"), ("torino", "Italy"),
    ("bologna", "Italy"), ("florence", "Italy"),
    ("united kingdom", "United Kingdom"), ("london", "United Kingdom"),
    ("england", "United Kingdom"), ("manchester", "United Kingdom"),
    ("edinburgh", "United Kingdom"), ("glasgow", "United Kingdom"), ("uk", "United Kingdom"),
    ("germany", "Germany"), ("deutschland", "Germany"), ("berlin", "Germany"),
    ("munich", "Germany"), ("münchen", "Germany"), ("frankfurt", "Germany"),
    ("hamburg", "Germany"), ("cologne", "Germany"), ("köln", "Germany"),
    ("france", "France"), ("paris", "France"), ("lyon", "France"), ("toulouse", "France"),
    ("netherlands", "Netherlands"), ("amsterdam", "Netherlands"),
    ("rotterdam", "Netherlands"), ("the hague", "Netherlands"), ("utrecht", "Netherlands"),
    ("spain", "Spain"), ("madrid", "Spain"), ("barcelona", "Spain"), ("valencia", "Spain"),
    ("ireland", "Ireland"), ("dublin", "Ireland"),
    ("switzerland", "Switzerland"), ("zurich", "Switzerland"), ("zürich", "Switzerland"),
    ("geneva", "Switzerland"), ("basel", "Switzerland"),
    ("sweden", "Sweden"), ("stockholm", "Sweden"), ("gothenburg", "Sweden"),
    ("belgium", "Belgium"), ("brussels", "Belgium"), ("antwerp", "Belgium"),
    ("austria", "Austria"), ("vienna", "Austria"), ("wien", "Austria"),
    ("portugal", "Portugal"), ("lisbon", "Portugal"), ("porto", "Portugal"),
    ("denmark", "Denmark"), ("copenhagen", "Denmark"),
    ("norway", "Norway"), ("oslo", "Norway"),
    ("finland", "Finland"), ("helsinki", "Finland"),
    ("poland", "Poland"), ("warsaw", "Poland"), ("kraków", "Poland"), ("krakow", "Poland"),
]

_REMOTE_RE = re.compile(r"\b(fully remote|remote.?first|work from anywhere|remote)\b", re.IGNORECASE)
_HYBRID_RE = re.compile(r"\bhybrid\b", re.IGNORECASE)


def _derive_country(location: str | None) -> str:
    """Best-effort country label from a free-text location string."""
    loc = (location or "").lower().strip()
    if not loc:
        return "Unknown"
    for keyword, country in _COUNTRY_LOOKUP:
        if keyword in loc:
            return country
    if re.search(r"\b(remote|anywhere|europe|european)\b", loc):
        return "Remote/EU"
    return location.strip()


def _derive_work_mode(location: str | None, description: str | None) -> str:
    """Classify work mode from location + the first part of the description."""
    text = f"{location or ''} {(description or '')[:600]}"
    if _HYBRID_RE.search(text):
        return "Hybrid"
    if _REMOTE_RE.search(text):
        return "Remote"
    return "On-site"


def _derive_sponsorship(job: dict) -> str:
    """Visa-sponsorship signal, derived from the eligibility visa gate."""
    gate = (job.get("visa_gate") or "").upper()
    if gate == "PASS":
        return "Likely"
    if gate == "FAIL":
        return "No"
    return "Unknown"


def export_jobs_master_xlsx(output_path: str | None = None) -> str:
    """
    Export the prioritized application dashboard to outputs/jobs_master.xlsx.

    Rows are ranked by score DESC then date-found DESC (see core.ranking) and
    color-coded by application-priority band. Overwrites the existing file each
    run so it remains a single source of truth for applications.

    Returns the absolute path of the written file.
    """
    from core.ranking import rank_jobs

    ranked = rank_jobs(database.get_all_scored_jobs_ranked())

    # INTERN_ONLY: restrict the dashboard to internship/graduate roles only, so
    # legacy non-intern rows already scored in the DB don't appear in the sheet.
    from core.prefilter import intern_only, is_intern_role
    if intern_only():
        ranked = [r for r in ranked if is_intern_role(r.get("title"))]

    # Enrich each row with the derived dashboard fields.
    for row in ranked:
        row["country"]            = _derive_country(row.get("location"))
        row["work_mode"]          = _derive_work_mode(row.get("location"), row.get("description"))
        row["sponsorship"]        = _derive_sponsorship(row)
        row["application_status"] = row.get("application_status") or "Not Started"
        row["notes"]              = row.get("application_notes") or ""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Jobs Master"
    _write_master_sheet(ws, ranked)

    # ── Additional sheets (best-effort — never break the core dashboard) ────────
    _safe_extra_sheets(wb, ranked)

    # ── Write file (default: outputs/jobs_master.xlsx — single source of truth) ─
    out = Path(output_path) if output_path else Path(config.OUTPUTS_DIR) / "jobs_master.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))

    # ── Timestamped archive copy (outputs/archive/jobs_master_YYYY_MM_DD_HHMM.xlsx)
    _archive_copy(out)

    return str(out.resolve())


def _style_header(ws, n_cols: int) -> None:
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 20
    ws.freeze_panes = "A2"


def _autofit(ws, n_cols: int, min_widths: list[int]) -> None:
    for col_idx in range(1, n_cols + 1):
        longest = max(
            (len(str(ws.cell(row=r, column=col_idx).value or ""))
             for r in range(1, ws.max_row + 1)),
            default=min_widths[col_idx - 1],
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(
            _MAX_COL_WIDTH, max(min_widths[col_idx - 1], longest + 2)
        )


def _write_master_sheet(ws, ranked: list[dict]) -> None:
    n = len(_MASTER_COLUMNS)
    ws.append([col[1] for col in _MASTER_COLUMNS])
    _style_header(ws, n)

    url_col = next(i + 1 for i, (k, _, _) in enumerate(_MASTER_COLUMNS) if k == "url")
    for row_data in ranked:
        ws.append([row_data.get(key) for key, _, _ in _MASTER_COLUMNS])
        row_n = ws.max_row
        fill = _PRIORITY_FILL.get(row_data.get("application_priority"))  # conditional by priority
        for col_idx in range(1, n + 1):
            cell = ws.cell(row=row_n, column=col_idx)
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(vertical="top")
        url_val = row_data.get("url")
        if url_val:
            url_cell = ws.cell(row=row_n, column=url_col)
            url_cell.hyperlink = url_val
            url_cell.font = Font(color="0563C1", underline="single")

    _autofit(ws, n, [c[2] for c in _MASTER_COLUMNS])
    ws.auto_filter.ref = ws.dimensions


# ── Secondary sheets ────────────────────────────────────────────────────────────

# Compact, action-oriented columns for the priority + applications sheets.
_PRIORITY_SHEET_COLUMNS: list[tuple[str, str, int]] = [
    ("rank",               "Rank",           6),
    ("total_score",        "Score",          7),
    ("recommendation",     "Recommendation", 18),
    ("company",            "Company",        24),
    ("title",              "Job Title",      42),
    ("location",           "Location",       22),
    ("sponsorship",        "Sponsorship",    12),
    ("application_status", "Status",         16),
    ("url",                "Job URL",        50),
]

_APPLICATIONS_COLUMNS: list[tuple[str, str, int]] = [
    ("job_id",       "Job ID",   8),
    ("total_score",  "Score",    7),
    ("status",       "Status",   16),
    ("company",      "Company",  24),
    ("title",        "Job Title", 42),
    ("location",     "Location", 22),
    ("notes",        "Notes",    30),
    ("last_updated", "Updated",  20),
]


def _write_simple_sheet(ws, columns, rows, url_key: str | None = "url") -> None:
    n = len(columns)
    ws.append([c[1] for c in columns])
    _style_header(ws, n)
    url_col = next((i + 1 for i, (k, _, _) in enumerate(columns) if k == url_key), None)
    for row_data in rows:
        ws.append([row_data.get(key) for key, _, _ in columns])
        if url_col:
            url_val = row_data.get(url_key)
            if url_val:
                cell = ws.cell(row=ws.max_row, column=url_col)
                cell.hyperlink = url_val
                cell.font = Font(color="0563C1", underline="single")
    _autofit(ws, n, [c[2] for c in columns])
    if ws.max_row >= 1:
        ws.auto_filter.ref = ws.dimensions


def _safe_extra_sheets(wb, ranked: list[dict]) -> None:
    """Add Apply Now / This Week / Applications / Analytics sheets.

    Wrapped so any data-layer hiccup degrades to the single-sheet dashboard
    rather than failing the whole export."""
    try:
        from core import analytics, tracker, database
        from core.recommendation_engine import recommend_all, sort_for_dashboard

        try:
            from core.profile_loader import load as _load_profile
            profile = _load_profile()
        except Exception:
            profile = None
        research_ids = database.get_researched_job_ids()

        enriched = sort_for_dashboard(
            recommend_all(ranked, profile=profile, research_job_ids=research_ids)
        )

        apply_now = [r for r in enriched if r.get("application_priority") == "Apply Now"]
        this_week = [r for r in enriched if r.get("application_priority") == "This Week"]
        _write_simple_sheet(wb.create_sheet("Apply Now"), _PRIORITY_SHEET_COLUMNS, apply_now)
        _write_simple_sheet(wb.create_sheet("This Week"), _PRIORITY_SHEET_COLUMNS, this_week)

        apps = tracker.list_applications()
        _write_simple_sheet(wb.create_sheet("Applications"), _APPLICATIONS_COLUMNS,
                            apps, url_key=None)

        _write_analytics_sheet(wb.create_sheet("Analytics"), analytics.summary())
    except Exception as exc:   # pragma: no cover - defensive
        print(f"  [exporter] extra sheets skipped: {exc}")


def _write_analytics_sheet(ws, summary: dict) -> None:
    from openpyxl.chart import BarChart, Reference

    am = summary["application_metrics"]
    cm = summary["conversion_metrics"]

    ws.append(["Career Analytics"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])

    # Summary cards (label / value)
    ws.append(["Metric", "Value"])
    hdr = ws.max_row
    for c in ws[hdr]:
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
    cards = [
        ("Jobs Found", am["jobs_found"]), ("Jobs Scored", am["jobs_scored"]),
        ("Companies Researched", am["researched"]), ("Documents", am["documents"]),
        ("Applied", am["applied"]), ("Assessments", am["assessments"]),
        ("Interviews", am["interviews"]), ("Final Rounds", am["final_rounds"]),
        ("Offers", am["offers"]), ("Rejections", am["rejections"]),
        ("Conversion Rate %", cm["conversion_rate"]),
        ("Interview Rate %", cm["interview_rate"]),
        ("Offer Rate %", cm["offer_rate"]),
    ]
    for label, value in cards:
        ws.append([label, value])

    # Funnel chart from the funnel rows.
    ws.append([])
    funnel_header_row = ws.max_row + 1
    ws.append(["Funnel Stage", "Count"])
    for c in ws[funnel_header_row]:
        c.fill, c.font = _HEADER_FILL, _HEADER_FONT
    funnel = summary["funnel_metrics"]
    for item in funnel:
        ws.append([item["stage"], item["count"]])
    last = ws.max_row

    chart = BarChart()
    chart.title = "Application Funnel"
    chart.type = "col"
    data = Reference(ws, min_col=2, min_row=funnel_header_row, max_row=last)
    cats = Reference(ws, min_col=1, min_row=funnel_header_row + 1, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 8, 16
    ws.add_chart(chart, "D3")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12


def _archive_copy(master_path: Path) -> Path:
    """Write a timestamped snapshot of the master dashboard to outputs/archive/."""
    import shutil

    stamp = datetime.now().strftime("%Y_%m_%d_%H%M")
    archive_dir = master_path.parent / "archive"
    archive_dir.mkdir(parents=True, exist_ok=True)
    dest = archive_dir / f"jobs_master_{stamp}.xlsx"
    shutil.copyfile(str(master_path), str(dest))
    return dest
